"""Reproduzierbare Auswertung der Messreihen des Reversionspendels.

Aufruf:
    python auswertung_reversionspendel.py "Rohdaten der Messreihen(2).xlsx"

Voraussetzung: openpyxl (pip install openpyxl).
Die Originalzeiten stehen in B:K, die Schwingungszahl n als Divisor der
Periodenformel in M, der Winkel in N und die reduzierte Länge in O.
Berechnete Zellwerte werden nicht als Eingaben verwendet.
"""

import argparse
import math
import re
import statistics

from openpyxl import load_workbook


DELTA_L = 0.001  # m, wie in der Dokumentation
N_MEASUREMENTS = 10


def inputs_from_row(sheet, row):
    """Lies ausschliesslich Originalzeiten und die Versuchsparameter.

    M enthält eine Formel für die Periodendauer. Ihr Ergebnis wird nicht
    verwendet; nur der Divisor ist die dokumentierte Schwingungszahl n.
    """
    times = [sheet.cell(row, col).value for col in range(2, 12)]
    formula = sheet.cell(row, 13).value
    match = re.fullmatch(r"=L" + str(row) + r"/(\d+)", str(formula))
    if not match or int(match.group(1)) <= 0:
        raise ValueError(f"{sheet.title}!M{row}: Schwingungszahl n nicht eindeutig")
    n = int(match.group(1))

    angle_text = str(sheet.cell(row, 14).value).strip()
    if not re.fullmatch(r"\d+(?:[.,]\d+)?\s*°", angle_text):
        raise ValueError(f"{sheet.title}!N{row}: Winkel nicht eindeutig")
    angle = float(angle_text.rstrip("°").strip().replace(",", "."))

    length = sheet.cell(row, 15).value
    if not isinstance(length, (int, float)) or length <= 0:
        raise ValueError(f"{sheet.title}!O{row}: Pendellänge nicht eindeutig")
    return times, n, float(length), angle


def evaluate_measurement(values, n, length, angle_degrees):
    """Berechne Einzelwert und Unsicherheit ohne Zwischenrundung."""
    if len(values) != N_MEASUREMENTS or not all(
        isinstance(value, (int, float)) for value in values
    ):
        raise ValueError("Eine Messreihe muss zehn numerische Einzelzeiten enthalten.")

    mean_time = statistics.mean(values)
    period = mean_time / n
    sample_s = statistics.stdev(values)  # Stichproben-Standardabweichung
    delta_t = sample_s / (n * math.sqrt(len(values)))
    g_calculated = 4 * math.pi**2 * length / period**2
    correction = 1 + math.radians(angle_degrees) ** 2 / 8
    g_corrected = g_calculated * correction
    delta_g = g_corrected * math.sqrt(
        (DELTA_L / length) ** 2 + (2 * delta_t / period) ** 2
    )
    return mean_time, period, sample_s, delta_t, g_calculated, g_corrected, delta_g


def main(path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    for sheet in workbook.worksheets:
        print(f"\n{sheet.title}")
        results = {"A": [], "B": []}
        series = None
        for row in range(1, sheet.max_row + 1):
            label = sheet.cell(row, 1).value
            if isinstance(label, str):
                series_match = re.match(r"\s*Reihe\s+(\d+)\b", label)
                if series_match:
                    series = int(series_match.group(1))
                blade_match = re.fullmatch(r"\s*Schneide\s+([AB])\s*", label)
                if not blade_match:
                    continue
                blade = blade_match.group(1)
                if series is None:
                    raise ValueError(f"{sheet.title}!A{row}: Reihe nicht zugeordnet")
                times, n, length, angle = inputs_from_row(sheet, row)
                mean_time, period, sample_s, delta_t, g_raw, g, delta_g = (
                    evaluate_measurement(times, n, length, angle)
                )
                weight = 1 / delta_g**2
                results[blade].append((g, weight))
                print(
                    f"  {series}{blade}: L={length:.3f} m, "
                    f"Winkel={angle:.2f}°, n={n}, "
                    f"t̄={mean_time:.8f} s, T={period:.8f} s, "
                    f"s={sample_s:.8f} s, "
                    f"ΔT={delta_t:.8f} s, g_ber={g_raw:.6f}, "
                    f"g_korr={g:.6f}, Δg={delta_g:.6f}, w={weight:.5f}"
                )

        if not results["A"] or len(results["A"]) != len(results["B"]):
            raise ValueError(f"{sheet.title}: unvollständige A/B-Messreihen")
        for blade in ("A", "B"):
            total_weight = sum(weight for _, weight in results[blade])
            weighted_g = sum(g * weight for g, weight in results[blade]) / total_weight
            uncertainty = math.sqrt(1 / total_weight)
            print(
                f"  Ergebnis {blade}: ({weighted_g:.3f} ± {uncertainty:.3f}) m/s² "
                f"(Σw={total_weight:.5f})"
            )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", help="Pfad zur Excel-Datei mit den Rohdaten")
    main(parser.parse_args().xlsx)
